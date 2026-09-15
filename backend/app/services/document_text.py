"""Plain-text extraction from uploaded documents (PDF / DOCX / PPTX).

Shared by resume import and experience extraction. Each function takes the raw
file bytes and returns the stripped text; parser errors propagate so the caller
can turn them into its own user-facing message.
"""
from io import BytesIO


def extract_pdf_text(raw: bytes) -> str:
    import pdfplumber

    chunks = []
    with pdfplumber.open(BytesIO(raw)) as pdf:
        for page in pdf.pages:
            chunks.append(page.extract_text() or "")
    return "\n".join(chunks).strip()


def extract_docx_text(raw: bytes) -> str:
    from docx import Document

    document = Document(BytesIO(raw))
    lines = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(line for line in lines if line.strip()).strip()


def extract_pptx_text(raw: bytes) -> str:
    """Slide by slide: text boxes, tables and grouped shapes, then speaker notes.

    Each slide with any text becomes a "[슬라이드 N]" block, so an image-only deck
    yields an empty string rather than a list of bare slide markers.
    """
    from pptx import Presentation
    from pptx.shapes.group import GroupShape

    def shape_lines(shape):
        if isinstance(shape, GroupShape):
            for child in shape.shapes:
                yield from shape_lines(child)
            return
        if getattr(shape, "has_table", False):
            for row in shape.table.rows:
                yield "\t".join(cell.text for cell in row.cells)
            return
        if getattr(shape, "has_text_frame", False):
            for paragraph in shape.text_frame.paragraphs:
                yield paragraph.text

    presentation = Presentation(BytesIO(raw))
    blocks = []
    for number, slide in enumerate(presentation.slides, start=1):
        lines = [line for shape in slide.shapes for line in shape_lines(shape)]
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame
            if notes is not None and notes.text.strip():
                lines.append(f"(발표자 노트) {notes.text}")
        # python-pptx reports a soft line break inside a paragraph as a vertical tab.
        lines = [line.replace("\v", "\n").strip() for line in lines]
        lines = [line for line in lines if line]
        if lines:
            blocks.append(f"[슬라이드 {number}]\n" + "\n".join(lines))
    return "\n\n".join(blocks).strip()
