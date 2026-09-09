"""Resume: an auto-drafted, user-editable summary of the user's timeline entries."""
import ipaddress
import re
import socket
from datetime import date, datetime
from io import BytesIO
from urllib.parse import urlparse
from urllib.request import urlopen

from fastapi import APIRouter, Depends, File, HTTPException, Response, UploadFile, status
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy.orm import Session

from app.database import get_db
from app.models import ActivityCategory, Resume, TimelineEntry, User
from app.schemas import (
    ResumeContent,
    ResumeImportContent,
    ResumeImportItem,
    ResumeImportOut,
    ResumeItem,
    ResumeOut,
    ResumeUpdate,
)
from app.security import get_current_user
from app.services.gemini import extract_resume_fields

router = APIRouter(prefix="/resume", tags=["resume"])

# Timeline categories -> the keys used inside resumes.content.
CATEGORY_KEYS: dict[ActivityCategory, str] = {
    ActivityCategory.EDUCATION: "education",
    ActivityCategory.CAREER: "career",
    ActivityCategory.ACTIVITY: "activity",
    ActivityCategory.CERTIFICATE: "certificate",
}

# The same mapping the other way round: resumes.content key -> timeline category.
SECTION_CATEGORIES: dict[str, ActivityCategory] = {key: category for category, key in CATEGORY_KEYS.items()}

# Section order + Korean headings, shared by the draft builder and the PDF renderer.
SECTIONS: list[tuple[str, str]] = [
    ("education", "학력"),
    ("career", "경력"),
    ("activity", "대외활동"),
    ("certificate", "자격증"),
]


def _get_resume(current_user: User, db: Session) -> Resume | None:
    return db.query(Resume).filter(Resume.user_id == current_user.id).first()


def _build_draft(current_user: User, db: Session) -> ResumeOut:
    """Build an unsaved resume from the user's profile + timeline entries.

    Nothing is written to the DB - this is only ever returned as a response.
    """
    entries = (
        db.query(TimelineEntry)
        .filter(TimelineEntry.user_id == current_user.id)
        .order_by(TimelineEntry.start_date.desc())
        .all()
    )
    grouped: dict[str, list[ResumeItem]] = {key: [] for key in CATEGORY_KEYS.values()}
    for entry in entries:
        key = CATEGORY_KEYS.get(entry.category)
        if key is None:
            continue
        grouped[key].append(
            ResumeItem(
                title=entry.title,
                start_date=entry.start_date,
                end_date=entry.end_date,
                timeline_entry_id=entry.id,
                section_label=None,
                description=None,
            )
        )
    return ResumeOut(
        user_id=current_user.id,
        name=current_user.name,
        email=current_user.email,
        phone=None,
        birth_date=None,
        photo_url=None,
        content=ResumeContent(**grouped),
        draft=True,
    )


def _current_resume(current_user: User, db: Session) -> ResumeOut:
    """The saved resume if there is one, otherwise a freshly built draft."""
    resume = _get_resume(current_user, db)
    if resume is None:
        return _build_draft(current_user, db)
    return ResumeOut.model_validate(resume)


@router.get("", response_model=ResumeOut)
def get_resume(current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    return _current_resume(current_user, db)


def _first_of_month(value: date | None) -> date | None:
    """Periods are year-month only, so the day is always stored as the 1st."""
    return value.replace(day=1) if value else None


def _normalize_periods(content: ResumeContent) -> None:
    for key in SECTION_CATEGORIES:
        for item in getattr(content, key):
            item.start_date = _first_of_month(item.start_date)
            item.end_date = _first_of_month(item.end_date)


def _sync_timeline_entries(content: ResumeContent, current_user: User, db: Session) -> None:
    """Mirror the resume's items into timeline_entries, filling in `timeline_entry_id`.

    An unlinked item creates a new entry; a linked one is updated in place, so the
    resume form doubles as the entry point for building a timeline.

    Nothing is ever deleted: dropping a row from the resume form must not destroy
    the timeline entry, because its items / sub-experiences (interview answers)
    hang off it. Unlinking is intentional - the entry stays in the timeline.
    """
    for key, category in SECTION_CATEGORIES.items():
        for item in getattr(content, key):
            if item.timeline_entry_id is not None:
                entry = db.get(TimelineEntry, item.timeline_entry_id)
                if entry is None or entry.user_id != current_user.id:
                    # Gone, or someone else's row: never write to it, and don't
                    # keep the dangling reference in this user's resume either.
                    item.timeline_entry_id = None
                    continue
                # `updated_at` is bumped by the column's onupdate when a field changes.
                entry.category = category
                entry.title = item.title
                entry.start_date = item.start_date
                entry.end_date = item.end_date
                continue

            entry = TimelineEntry(
                user_id=current_user.id,
                category=category,
                title=item.title,
                start_date=item.start_date,
                end_date=item.end_date,
            )
            db.add(entry)
            db.flush()  # assign the PK so it can be stored back into the resume
            item.timeline_entry_id = entry.id


@router.put("", response_model=ResumeOut)
def upsert_resume(
    payload: ResumeUpdate,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    phone = payload.phone.strip() if payload.phone else None
    photo_url = payload.photo_url.strip() if payload.photo_url else None
    _normalize_periods(payload.content)
    # Fills in timeline_entry_id on the items, so the saved content carries the links.
    _sync_timeline_entries(payload.content, current_user, db)
    content = payload.content.model_dump(mode="json")

    resume = _get_resume(current_user, db)
    if resume is None:
        resume = Resume(
            user_id=current_user.id,
            name=payload.name,
            email=payload.email,
            phone=phone,
            birth_date=payload.birth_date,
            photo_url=photo_url,
            content=content,
        )
        db.add(resume)
    else:
        resume.name = payload.name
        resume.email = payload.email
        resume.phone = phone
        resume.birth_date = payload.birth_date
        resume.photo_url = photo_url
        resume.content = content
    db.commit()
    db.refresh(resume)
    return resume


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

# reportlab's built-in fonts are Latin-only; the Adobe CJK CID fonts ship with
# reportlab itself, so Hangul renders without any extra font file or native dep.
# Neither has a bold cut, so the gothic face stands in for emphasis against the
# lighter myeongjo body text.
_KO_BODY_FONT = "HYSMyeongJo-Medium"
_KO_BOLD_FONT = "HYGothic-Medium"
_FALLBACK_FONT = "Helvetica"
_FALLBACK_BOLD_FONT = "Helvetica-Bold"

# Cap on the profile photo we are willing to pull in (bytes) and how long to wait.
_PHOTO_MAX_BYTES = 5 * 1024 * 1024
_PHOTO_TIMEOUT_SECONDS = 5


def _resolve_font(name: str, fallback: str) -> str:
    try:
        pdfmetrics.getFont(name)
        return name
    except KeyError:
        pass
    try:
        pdfmetrics.registerFont(UnicodeCIDFont(name))
        return name
    except Exception:
        # Never fail the download over a font: Latin text still renders fine.
        return fallback


def _format_year_month(value: str | None) -> str:
    """'2021-03-01' -> '2021. 03.' - periods are year-month only, never the day."""
    if not value:
        return ""
    parts = value.split("-")
    return f"{parts[0]}. {parts[1]}." if len(parts) >= 2 else value


def _format_period(item: dict) -> str:
    start = _format_year_month(item.get("start_date"))
    end = _format_year_month(item.get("end_date")) or "현재"
    return f"{start} – {end}" if start else end


def _korean_age(birth: date, today: date) -> int:
    """만 나이: full years elapsed, so the birthday must already have passed."""
    age = today.year - birth.year
    if (today.month, today.day) < (birth.month, birth.day):
        age -= 1
    return age


def _format_birth_date(value: str | None) -> str | None:
    """'1997-10-25' -> '1997. 10. 25. (만 28세)'."""
    if not value:
        return None
    try:
        birth = date.fromisoformat(value)
    except ValueError:
        return value
    formatted = f"{birth.year}. {birth.month}. {birth.day}."
    age = _korean_age(birth, date.today())
    return f"{formatted} (만 {age}세)" if age >= 0 else formatted


def _is_public_http_url(url: str) -> bool:
    """Only plain http(s) to a public address - the server does the fetching."""
    parsed = urlparse(url)
    if parsed.scheme not in ("http", "https") or not parsed.hostname:
        return False
    try:
        infos = socket.getaddrinfo(parsed.hostname, None)
    except socket.gaierror:
        return False
    for info in infos:
        try:
            ip = ipaddress.ip_address(info[4][0])
        except ValueError:
            return False
        if ip.is_private or ip.is_loopback or ip.is_link_local or ip.is_reserved:
            return False
    return True


def _fetch_photo(url: str | None) -> bytes | None:
    """Best-effort profile photo bytes. Any problem just means "no photo" - never a 500."""
    if not url or not _is_public_http_url(url):
        return None
    try:
        with urlopen(url, timeout=_PHOTO_TIMEOUT_SECONDS) as response:  # noqa: S310 - scheme checked above
            raw = response.read(_PHOTO_MAX_BYTES + 1)
        if not raw or len(raw) > _PHOTO_MAX_BYTES:
            return None
        ImageReader(BytesIO(raw)).getSize()  # forces a decode, so a broken file fails here
        return raw
    except Exception:
        return None


def _grouped_sections(content: dict) -> list[tuple[str, list[dict]]]:
    """Group items by their display heading, in first-appearance order.

    An item's own `section_label` wins; otherwise it falls back to the Korean
    label of the category it sits under.
    """
    groups: dict[str, list[dict]] = {}
    for key, default_label in SECTIONS:
        for item in content.get(key) or []:
            label = (item.get("section_label") or "").strip() or default_label
            groups.setdefault(label, []).append(item)
    return list(groups.items())


# Imported/pasted text often already carries its own bullet marker; strip it so
# the renderer doesn't produce "- - text".
_BULLET_PREFIX_RE = re.compile(r"^[-*\u2022\u00b7\u25cf\u25aa]+\s*")


def _bullets(item: dict) -> list[str]:
    lines = []
    for raw in (item.get("description") or "").splitlines():
        line = _BULLET_PREFIX_RE.sub("", raw.strip()).strip()
        if line:
            lines.append(line)
    return lines


def _render_pdf(data: ResumeOut, author_name: str) -> bytes:
    body_font = _resolve_font(_KO_BODY_FONT, _FALLBACK_FONT)
    bold_font = _resolve_font(_KO_BOLD_FONT, _FALLBACK_BOLD_FONT)
    content = data.content.model_dump(mode="json")

    name_style = ParagraphStyle("cbName", fontName=bold_font, fontSize=20, leading=26, spaceAfter=6)
    meta_style = ParagraphStyle("cbMeta", fontName=body_font, fontSize=10, leading=16, textColor=colors.HexColor("#444444"))
    heading_style = ParagraphStyle("cbHeading", fontName=bold_font, fontSize=13, leading=18)
    period_style = ParagraphStyle("cbPeriod", fontName=body_font, fontSize=9.5, leading=14, textColor=colors.HexColor("#555555"))
    item_title_style = ParagraphStyle("cbItemTitle", fontName=bold_font, fontSize=11, leading=16)
    bullet_style = ParagraphStyle(
        "cbBullet", fontName=body_font, fontSize=9.5, leading=15, leftIndent=9, firstLineIndent=-9
    )
    empty_style = ParagraphStyle("cbEmpty", fontName=body_font, fontSize=9.5, leading=14, textColor=colors.HexColor("#888888"))
    confirm_style = ParagraphStyle("cbConfirm", fontName=body_font, fontSize=10, leading=16, alignment=TA_CENTER)
    author_style = ParagraphStyle("cbAuthor", fontName=bold_font, fontSize=10, leading=16, alignment=TA_CENTER)

    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=20 * mm,
        rightMargin=20 * mm,
        topMargin=20 * mm,
        bottomMargin=20 * mm,
        title="이력서",
        author=author_name,
    )
    table_width = doc.width

    # ---- header: name + personal details on the left, photo on the right ----
    details = [Paragraph(data.name, name_style)]
    birth = _format_birth_date(data.birth_date.isoformat() if data.birth_date else None)
    for label, value in (("생년월일", birth), ("전화번호", data.phone), ("이메일", data.email)):
        if value:
            details.append(Paragraph(f"{label} &nbsp;&nbsp; {value}", meta_style))

    photo = _fetch_photo(data.photo_url)
    if photo is None:
        flow: list = list(details)
    else:
        photo_w = 28 * mm
        # Separate streams: ImageReader and Image each consume the buffer they get.
        img_w, img_h = ImageReader(BytesIO(photo)).getSize()
        photo_h = photo_w * img_h / img_w if img_w else photo_w
        header = Table(
            [[details, Image(BytesIO(photo), width=photo_w, height=photo_h)]],
            colWidths=[table_width - photo_w - 6 * mm, photo_w + 6 * mm],
        )
        header.setStyle(
            TableStyle(
                [
                    ("VALIGN", (0, 0), (0, 0), "TOP"),
                    ("VALIGN", (1, 0), (1, 0), "TOP"),
                    ("ALIGN", (1, 0), (1, 0), "RIGHT"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 0),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 0),
                    ("TOPPADDING", (0, 0), (-1, -1), 0),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
                ]
            )
        )
        flow = [header]

    flow.append(Spacer(1, 10))

    # ---- sections: heading, then "period | title + bullets" rows ----
    groups = _grouped_sections(content)
    if not groups:
        flow.append(Paragraph("등록된 항목이 없어요.", empty_style))

    for label, items in groups:
        heading = Table([[Paragraph(label, heading_style)]], colWidths=[table_width])
        heading.setStyle(
            TableStyle(
                [
                    ("LEFTPADDING", (0, 0), (-1, -1), 0),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 0),
                    ("TOPPADDING", (0, 0), (-1, -1), 12),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                    ("LINEBELOW", (0, 0), (-1, -1), 0.8, colors.HexColor("#333333")),
                ]
            )
        )
        flow.append(heading)

        rows = []
        for item in items:
            cell = [Paragraph(item.get("title", ""), item_title_style)]
            cell += [Paragraph(f"- {line}", bullet_style) for line in _bullets(item)]
            rows.append([Paragraph(_format_period(item), period_style), cell])
        table = Table(rows, colWidths=[table_width * 0.3, table_width * 0.7])
        table.setStyle(
            TableStyle(
                [
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("TOPPADDING", (0, 0), (-1, -1), 7),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
                    ("LEFTPADDING", (0, 0), (-1, -1), 0),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 0),
                ]
            )
        )
        flow.append(table)

    # ---- closing statement ----
    flow.append(Spacer(1, 26))
    flow.append(Paragraph("위의 모든 기재사항은 사실과 다름없음을 확인합니다.", confirm_style))
    flow.append(Spacer(1, 8))
    flow.append(Paragraph(f"작성자: {author_name}", author_style))

    doc.build(flow)
    return buffer.getvalue()


@router.get("/pdf", responses={200: {"content": {"application/pdf": {}}, "description": "이력서 PDF"}})
def download_resume_pdf(
    current_user: User = Depends(get_current_user), db: Session = Depends(get_db)
):
    data = _current_resume(current_user, db)
    # The closing line is signed by the logged-in account, not the editable header name.
    pdf = _render_pdf(data, current_user.name)
    return Response(
        content=pdf,
        media_type="application/pdf",
        headers={"Content-Disposition": 'attachment; filename="resume.pdf"'},
    )


# ---------------------------------------------------------------------------
# Import: pre-fill the form from an uploaded resume file
# ---------------------------------------------------------------------------

_IMPORT_MAX_BYTES = 10 * 1024 * 1024
_TITLE_MAX = 50
_SECTION_LABEL_MAX = 50
_DESCRIPTION_MAX = 2000

_NO_AI_WARNING = (
    "AI 키가 없어 자동 정리는 어렵지만 텍스트는 추출했습니다. "
    "설명란의 원문을 보고 직접 항목을 정리한 뒤 저장해주세요."
)


def _extract_pdf_text(raw: bytes) -> str:
    import pdfplumber

    chunks = []
    with pdfplumber.open(BytesIO(raw)) as pdf:
        for page in pdf.pages:
            chunks.append(page.extract_text() or "")
    return "\n".join(chunks).strip()


def _extract_docx_text(raw: bytes) -> str:
    from docx import Document

    document = Document(BytesIO(raw))
    lines = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(line for line in lines if line.strip()).strip()


def _clip(value, limit: int) -> str | None:
    """Trim AI/raw output to what the resume schema accepts."""
    if not isinstance(value, str):
        return None
    trimmed = value.strip()
    return trimmed[:limit] if trimmed else None


def _parse_month(value) -> date | None:
    """Accepts 'YYYY-MM' or 'YYYY-MM-DD'; always returns the 1st of that month."""
    if not isinstance(value, str):
        return None
    text = value.strip()
    for fmt in ("%Y-%m", "%Y-%m-%d", "%Y.%m", "%Y/%m"):
        try:
            return datetime.strptime(text, fmt).date().replace(day=1)
        except ValueError:
            continue
    return None


def _parse_birth_date(value) -> date | None:
    if not isinstance(value, str):
        return None
    try:
        return date.fromisoformat(value.strip())
    except ValueError:
        return None


def _to_import_items(raw_items) -> list[ResumeImportItem]:
    items: list[ResumeImportItem] = []
    if not isinstance(raw_items, list):
        return items
    for raw in raw_items:
        if not isinstance(raw, dict):
            continue
        items.append(
            ResumeImportItem(
                title=_clip(raw.get("title"), _TITLE_MAX) or "",
                start_date=_parse_month(raw.get("start_date")),
                end_date=_parse_month(raw.get("end_date")),
                section_label=_clip(raw.get("section_label"), _SECTION_LABEL_MAX),
                description=_clip(raw.get("description"), _DESCRIPTION_MAX),
            )
        )
    return items


@router.post("/import", response_model=ResumeImportOut)
def import_resume(
    file: UploadFile = File(...),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """Turn an uploaded PDF/DOCX into an unsaved draft that pre-fills the form.

    Nothing is written here - no resume row, no timeline entries. The user
    reviews the parsed values and presses save, which runs the normal upsert.
    """
    filename = file.filename or ""
    suffix = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if suffix not in ("pdf", "docx"):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="PDF 또는 DOCX 파일만 올릴 수 있어요.",
        )

    raw = file.file.read(_IMPORT_MAX_BYTES + 1)
    if not raw:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="빈 파일이에요.")
    if len(raw) > _IMPORT_MAX_BYTES:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST, detail="파일 용량은 10MB 이하여야 해요."
        )

    try:
        text = _extract_pdf_text(raw) if suffix == "pdf" else _extract_docx_text(raw)
    except Exception:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="파일에서 텍스트를 읽지 못했어요. 다른 파일로 다시 시도해주세요.",
        )

    if not text:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="파일에서 글자를 찾지 못했어요. 스캔 이미지 PDF는 인식할 수 없어요.",
        )

    parsed = extract_resume_fields(text)

    if parsed is None:
        # Fallback: hand the raw text back in one item's description so the user
        # can reorganise it manually instead of losing the upload entirely.
        stem = filename.rsplit(".", 1)[0] or "첨부한 이력서"
        content = ResumeImportContent(
            education=[
                ResumeImportItem(
                    title=stem[:_TITLE_MAX],
                    section_label="첨부한 이력서 원문",
                    description=text[:_DESCRIPTION_MAX],
                )
            ]
        )
        return ResumeImportOut(
            name=current_user.name,
            email=current_user.email,
            content=content,
            warning=_NO_AI_WARNING,
        )

    raw_content = parsed.get("content") or {}
    content = ResumeImportContent(
        **{key: _to_import_items(raw_content.get(key)) for key in SECTION_CATEGORIES}
    )
    return ResumeImportOut(
        name=_clip(parsed.get("name"), 100) or current_user.name,
        email=current_user.email,
        phone=_clip(parsed.get("phone"), 50),
        birth_date=_parse_birth_date(parsed.get("birth_date")),
        content=content,
    )

